"""The attendance grid saves one change at a time, so the server must treat a save as PARTIAL: only the students named
change, everyone else on that date is left exactly as they were."""
from helpers import Api, make_account


def url(c, path=""):
    return f"/api/teach/offerings/{c['oid']}/attendance{path}"


def second_student(db, c):
    s3 = make_account(db, "s3@example.com", "student", "S-3")
    assert c["admin"].post(f"/api/sections/{c['sec_a']['id']}/members", {"student_id": str(s3.id)}).status_code in (200, 201, 204)
    return s3


def marks_on(c, date):
    data = c["fac"].get(url(c) + f"?section_id={c['sec_a']['id']}").json()
    day = next((s for s in data["sessions"] if s["date"] == date), None)
    return day["marks"] if day else None


def save(c, date, marks, period="midterm", who="fac"):
    return c[who].put(url(c), {"section_id": c["sec_a"]["id"], "session_date": date, "period": period, "marks": marks})


def test_a_save_only_changes_the_students_it_names(db, graded):
    c = graded
    s1, s3 = c["students"][0], second_student(db, c)
    assert save(c, "2026-07-06", {str(s1.id): "present", str(s3.id): "late"}).status_code == 204
    assert marks_on(c, "2026-07-06") == {str(s1.id): "present", str(s3.id): "late"}
    assert save(c, "2026-07-06", {str(s1.id): "absent"}).status_code == 204                 # one cell edited in the grid
    assert marks_on(c, "2026-07-06") == {str(s1.id): "absent", str(s3.id): "late"}           # the other student is untouched
    assert save(c, "2026-07-06", {str(s1.id): None}).status_code == 204                       # a cleared cell
    assert marks_on(c, "2026-07-06") == {str(s3.id): "late"}
    assert save(c, "2026-07-06", {str(s1.id): "present"}).status_code == 204 and len(marks_on(c, "2026-07-06")) == 2


def test_each_date_is_its_own_session_and_a_dates_period_cannot_change_once_recorded(db, graded):
    c = graded
    s1 = c["students"][0]
    assert save(c, "2026-07-06", {str(s1.id): "present"}, period="midterm").status_code == 204
    assert save(c, "2026-07-08", {str(s1.id): "late"}, period="finals").status_code == 204
    assert marks_on(c, "2026-07-06") == {str(s1.id): "present"} and marks_on(c, "2026-07-08") == {str(s1.id): "late"}
    clash = save(c, "2026-07-06", {str(s1.id): "absent"}, period="finals")
    assert clash.status_code == 409 and clash.json()["error"]["code"] == "period_locked"
    assert marks_on(c, "2026-07-06") == {str(s1.id): "present"}                                 # the clash changed nothing


def test_only_students_of_the_section_can_be_marked_and_only_by_their_teacher(db, graded):
    c = graded
    s1, outsider = c["students"][0], c["students"][1]                                           # s2 is in section 1B, not 1A
    stray = save(c, "2026-07-06", {str(outsider.id): "present"})
    assert stray.status_code == 422 and stray.json()["error"]["code"] == "student_invalid"
    other = make_account(db, "attendance-other@example.com", "faculty")
    assert Api(other.email).put(url(c), {"section_id": c["sec_a"]["id"], "session_date": "2026-07-06", "period": "midterm", "marks": {str(s1.id): "present"}}).status_code == 404
    assert save(c, "2026-07-06", {str(s1.id): "present"}, who="st1").status_code == 403
    assert save(c, "2026-07-06", {str(s1.id): "present"}, who="admin").status_code == 403
    assert marks_on(c, "2026-07-06") is None
    assert c["admin"].post(f"/api/terms/{c['term']['id']}/close").status_code in (200, 204)
    closed = save(c, "2026-07-06", {str(s1.id): "present"})
    assert closed.status_code == 409 and closed.json()["error"]["code"] == "term_closed"


# ---------------- export ----------------
from io import BytesIO
from urllib.parse import urlencode

from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select

from app.features.accounts.models import Account, AuditEvent


def export(c, fmt="xlsx", who="fac", **query):
    params = {"format": fmt, "section_id": c["sec_a"]["id"], **query}
    return c[who].get(f"/api/teach/offerings/{c['oid']}/exports/attendance?{urlencode(params)}")


def sheet_rows(response):
    sheet = load_workbook(BytesIO(response.content)).active
    return sheet.title, [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]


def record_three_days(db, c):
    s1, s3 = c["students"][0], second_student(db, c)
    save(c, "2026-07-06", {str(s1.id): "present", str(s3.id): "late"}, period="midterm")
    save(c, "2026-07-08", {str(s1.id): "absent"}, period="midterm")                       # s3 not marked that day
    save(c, "2026-09-14", {str(s1.id): "excused", str(s3.id): "present"}, period="finals")
    return s1, s3


def test_the_attendance_export_lists_recorded_dates_marks_and_totals(db, graded):
    c = graded
    s1, s3 = record_three_days(db, c)
    response = export(c)
    assert response.status_code == 200 and response.headers["content-disposition"].endswith('.xlsx"')
    assert response.headers["cache-control"] == "no-store" and response.headers["x-content-type-options"] == "nosniff"
    title, rows = sheet_rows(response)
    assert title == "Attendance"
    header = next(r for r in rows if r and r[0] == "Student no.")
    assert header[2:5] == ["Mon Jul 06 · Mid", "Wed Jul 08 · Mid", "Mon Sep 14 · Fin"] and header[5:] == ["Present", "Late", "Absent", "Excused", "Not marked"]
    by_name = {r[1]: r for r in rows if r and r[1] in (s1.display_name, s3.display_name)}
    assert by_name[s1.display_name][2:5] == ["P", "A", "E"] and by_name[s1.display_name][5:] == [1, 0, 1, 1, 0]
    assert by_name[s3.display_name][2:5] == ["L", None, "P"] or by_name[s3.display_name][2:5] == ["L", "", "P"]
    assert by_name[s3.display_name][5:] == [1, 1, 0, 0, 1]                                  # never marked on Jul 8: blank, not absent
    flat = " ".join(str(v) for r in rows for v in r if v)
    assert "not an official institutional form" in flat and "Recorded days" in flat


def test_the_export_can_be_limited_to_a_range_and_refuses_an_empty_or_invalid_one(db, graded):
    c = graded
    record_three_days(db, c)
    _, rows = sheet_rows(export(c, start="2026-07-01", end="2026-07-31"))
    header = next(r for r in rows if r and r[0] == "Student no.")
    assert header[2:4] == ["Mon Jul 06", "Wed Jul 08"] and "Present" in header and len(header) == 2 + 2 + 5     # one period: no suffix
    only = export(c, start="2026-09-01")
    assert only.status_code == 200 and "Sep 14" in str(sheet_rows(only)[1])
    empty = export(c, start="2027-01-01", end="2027-01-31")
    assert empty.status_code == 422 and empty.json()["error"]["code"] == "no_attendance"
    assert export(c, start="2026-08-01", end="2026-07-01").json()["error"]["code"] == "range_invalid"
    assert c["fac"].get(f"/api/teach/offerings/{c['oid']}/exports/attendance?" + urlencode({"format": "xlsx", "section_id": c["sec_b"]["id"]})).status_code == 422   # 1B does not take this subject here


def test_names_that_start_like_a_formula_are_neutralised_in_the_file(db, graded):
    c = graded
    s1 = c["students"][0]
    db.get(Account, s1.id).display_name = "=HYPERLINK(\"http://evil\",\"x\")"
    db.commit()
    save(c, "2026-07-06", {str(s1.id): "present"})
    _, rows = sheet_rows(export(c))
    assert any(r[1] == "'=HYPERLINK(\"http://evil\",\"x\")" for r in rows if r and len(r) > 1)


def test_the_pdf_splits_a_long_range_across_tables_and_is_audited(db, graded):
    c = graded
    s1 = c["students"][0]
    for day in range(1, 26):
        assert save(c, f"2026-07-{day:02d}", {str(s1.id): "present"}).status_code == 204
    response = export(c, "pdf")
    assert response.status_code == 200 and response.content.startswith(b"%PDF") and response.headers["content-disposition"].endswith('.pdf"')
    reader = PdfReader(BytesIO(response.content))
    text = " ".join(page.extract_text() for page in reader.pages)
    assert len(reader.pages) >= 2 and "Dates 1 to 18 of 25" in text and "Dates 19 to 25 of 25" in text and s1.display_name in text
    event = db.scalars(select(AuditEvent).where(AuditEvent.action == "attendance.exported")).one()
    assert event.details["dates"] == 25 and event.details["format"] == "pdf" and s1.display_name not in str(event.details)


def test_only_the_teacher_can_export_attendance_and_a_closed_term_still_can(db, graded):
    c = graded
    record_three_days(db, c)
    from helpers import Api
    other = make_account(db, "attendance-export@example.com", "faculty")
    assert Api(other.email).get(f"/api/teach/offerings/{c['oid']}/exports/attendance?" + urlencode({"format": "xlsx", "section_id": c["sec_a"]["id"]})).status_code == 404
    assert export(c, who="st1").status_code == 403 and export(c, who="admin").status_code == 403
    assert c["admin"].post(f"/api/terms/{c['term']['id']}/close").status_code in (200, 204)
    assert export(c).status_code == 200 and export(c, "pdf").status_code == 200

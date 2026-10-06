import uuid
from io import BytesIO

import pytest
from helpers import Api, make_account
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select
from test_activities import score
from test_grading import book, full_term, g, publish_grades

from app.features.accounts.models import Account
from app.features.exports import render


def export(c, who="fac", **q):
    q = {"format": "xlsx", "period": "midterm", **q}
    query = "&".join(f"{k}={v}" for k, v in q.items())
    return c[who].get(g(c, f"/exports/grades?{query}"))


def sheet_rows(response):
    """(banner_or_None, meta dict, header, data rows) of the first worksheet."""
    ws = load_workbook(BytesIO(response.content)).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    start = next(i for i, r in enumerate(rows) if r[0] in ("Student no.", "Rank"))
    banner = rows[0][0] if rows[0][0] and str(rows[0][0]).startswith("PREVIEW") else None
    meta = {r[0]: r[1] for r in rows[:start] if r[0] and r[1] is not None}
    return banner, meta, rows[start], rows[start + 1:]


def pdf_pages(response):
    return [p.extract_text() for p in PdfReader(BytesIO(response.content)).pages]


def test_published_export_is_what_students_can_see_and_nothing_more(graded):
    full_term(graded)
    publish_grades(graded, "midterm")                                  # s1 released, s2 pending
    response = export(graded, basis="published")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in response.headers["content-disposition"] and response.headers["cache-control"] == "no-store"
    banner, meta, header, rows = sheet_rows(response)
    assert banner is None                                              # released data is not a preview
    assert meta["Period"] == "Midterm" and meta["Basis"].startswith("Published")
    assert meta["Subject"].count("-") >= 1 and meta["Students"] == "2"
    assert header[:3] == ["Student no.", "Name", "Section"] and header[-3:] == ["Grade", "Remark", "Status"]
    by_name = {r[1]: r for r in rows}
    s1, s2 = by_name["s1"], by_name["s2"]
    values = dict(zip(header, s1, strict=True))
    assert values["Grade"] == 81.0 and values["Remark"] == "passed" and values["Section"] == "1A"
    assert [values[h] for h in header if h.endswith("(%)")] == [80.0, 90.0, 70.0, 80.0]
    assert values["Status"].startswith("Published release 1")
    assert s2[header.index("Grade")] is None and s2[-1] == "Not published"     # no release -> no number


def test_published_export_keeps_the_release_even_after_working_data_changes(graded):
    ids = full_term(graded)
    publish_grades(graded, "midterm")
    assert score(graded, ids["exam_m"], graded["students"][0], 44).status_code == 200      # working value becomes 85
    _, _, header, rows = sheet_rows(export(graded, basis="published"))
    s1 = dict(zip(header, next(r for r in rows if r[1] == "s1"), strict=True))
    assert s1["Grade"] == 81.0 and "review pending" in s1["Status"]
    banner, _, header, rows = sheet_rows(export(graded, basis="working"))
    assert banner and dict(zip(header, next(r for r in rows if r[1] == "s1"), strict=True))["Grade"] == 85.0
    publish_grades(graded, "midterm")
    _, _, header, rows = sheet_rows(export(graded, basis="published"))
    s1 = dict(zip(header, next(r for r in rows if r[1] == "s1"), strict=True))
    assert s1["Grade"] == 85.0 and s1["Status"].startswith("Published release 2") and "review" not in s1["Status"]


def test_working_export_is_visibly_a_preview_everywhere(graded):
    full_term(graded)
    banner, meta, header, rows = sheet_rows(export(graded, basis="working"))
    assert "PREVIEW - NOT PUBLISHED" in banner and meta["Basis"].startswith("Working")
    by_name = {r[1]: dict(zip(header, r, strict=True)) for r in rows}
    assert by_name["s1"]["Grade"] == 81.0 and by_name["s1"]["Status"] == "Working value (preview)"
    assert by_name["s2"]["Grade"] is None and by_name["s2"]["Status"].startswith("Pending:")   # nothing guessed
    response = export(graded, basis="working", format="pdf")
    assert response.headers["content-type"] == "application/pdf" and "PREVIEW" in response.headers["content-disposition"]
    pages = pdf_pages(response)
    assert all("PREVIEW - NOT PUBLISHED" in p for p in pages)
    assert "s1" in pages[0] and "81.00" in pages[0] and "not an official institutional form" in pages[0]
    published_pdf = pdf_pages(export(graded, basis="published", format="pdf"))
    assert not any("PREVIEW" in p for p in published_pdf) and "Not published" in published_pdf[0]


def test_every_pdf_page_of_a_long_preview_carries_the_banner():
    table = {"title": "T", "banner": "PREVIEW - NOT PUBLISHED. test", "meta": [("Subject", "X")],
             "notes": ["Not official."], "columns": ["Student no.", "Name", "Section", "Quiz (%)", "Grade", "Remark", "Status"],
             "rows": [[f"S-{n}", f"Student {n}", "1A", 80.0, 80.0, "passed", "Working value (preview)"] for n in range(150)]}
    pages = [p.extract_text() for p in PdfReader(BytesIO(render.to_pdf(table))).pages]
    assert len(pages) >= 3 and all("PREVIEW - NOT PUBLISHED" in p for p in pages)
    assert "Student 149" in pages[-1]


def test_section_and_period_selection(graded):
    full_term(graded)
    publish_grades(graded, "midterm")
    _, meta, _, rows = sheet_rows(export(graded, basis="published", section_id=graded["sec_b"]["id"]))
    assert [r[1] for r in rows] == ["s2"] and meta["Sections"] == "1B" and meta["Students"] == "1"
    _, _, _, rows = sheet_rows(export(graded, basis="published", section_id=graded["sec_a"]["id"]))
    assert [r[1] for r in rows] == ["s1"]
    foreign = export(graded, section_id=uuid.uuid4())
    assert foreign.status_code == 422 and foreign.json()["error"]["code"] == "section_invalid"
    publish_grades(graded, "finals")
    publish_grades(graded, "course")
    _, meta, header, rows = sheet_rows(export(graded, period="course", basis="published"))
    s1 = dict(zip(header, next(r for r in rows if r[1] == "s1"), strict=True))
    assert meta["Period"].startswith("Course") and (s1["Midterm grade (used for this Course release)"], s1["Finals grade (used for this Course release)"], s1["Grade"]) == (81.0, 93.0, 87.0)
    _, _, header, rows = sheet_rows(export(graded, period="finals", basis="published"))
    assert dict(zip(header, next(r for r in rows if r[1] == "s1"), strict=True))["Grade"] == 93.0
    assert export(graded, period="weekly").status_code == 422
    assert export(graded, basis="secret").status_code == 422


def test_only_the_teaching_owner_can_export_and_exports_are_audited(db, graded):
    full_term(graded)
    for who in ("admin", "st1", "st2"):
        assert export(graded, who=who).status_code == 403
    stranger = Api(make_account(db, "other@example.com", "faculty").email)
    assert stranger.get(g(graded, "/exports/grades?format=pdf&period=midterm")).status_code == 404
    from fastapi.testclient import TestClient

    from app.main import app
    assert TestClient(app).get(g(graded, "/exports/grades?format=pdf&period=midterm")).status_code == 401
    assert export(graded, format="pdf").status_code == 200 and export(graded, basis="working").status_code == 200
    events = [e for e in graded["admin"].get("/api/audit").json() if e["action"] == "grades.exported"]
    assert len(events) == 2 and {e["details"]["format"] for e in events} == {"pdf", "xlsx"}
    assert {e["details"]["basis"] for e in events} == {"published", "working"}


def test_closed_terms_can_still_export_their_history(graded):
    full_term(graded)
    publish_grades(graded, "midterm")
    graded["admin"].post(f"/api/terms/{graded['term']['id']}/close")
    assert export(graded).status_code == 200 and export(graded, format="pdf", basis="working").status_code == 200


def test_spreadsheet_formulas_in_names_are_neutralised(db, graded):
    full_term(graded)
    account = db.scalar(select(Account).where(Account.email == "s1@example.com"))
    account.display_name = '=HYPERLINK("http://evil.example","click")'
    account.student_number = "+1-555"
    db.commit()
    _, _, _, rows = sheet_rows(export(graded, basis="working"))
    first = next(r for r in rows if "HYPERLINK" in str(r[1]))
    assert first[1].startswith("'=") and first[0].startswith("'+")
    ws = load_workbook(BytesIO(export(graded, basis="working").content)).active
    assert not any(c.data_type == "f" for row in ws.iter_rows() for c in row)       # no formula cells at all


@pytest.mark.parametrize("name,expected", [("a b/c", "a-b-c"), ("../../x", "x"), ("", "export")])
def test_download_names_are_safe(name, expected):
    assert render.safe_filename(name) == expected


def test_published_metadata_comes_from_the_release_not_from_todays_policy(db, graded):
    from helpers import POLICY, publish_syllabus
    full_term(graded)
    publish_grades(graded, "midterm")                                     # released under policy version 1, passing 75
    publish_syllabus(graded, policy={**POLICY, "passing": 90})            # policy version 2; the release is untouched
    _, meta, header, rows = sheet_rows(export(graded, basis="published"))
    assert "version 1; passing grade 75" in meta["Grading policy"] and meta["Grading policy"].endswith("(as released)")
    assert "Policy version" not in header                                 # one version, so one line is enough
    s1 = dict(zip(header, next(r for r in rows if r[1] == "s1"), strict=True))
    assert s1["Remark"] == "passed" and s1["Grade"] == 81.0               # still what the student saw
    _, meta, _, _ = sheet_rows(export(graded, basis="working"))
    assert "version 2; passing grade 90" in meta["Grading policy"] and meta["Grading policy"].endswith("(current)")


def test_releases_from_different_policy_versions_are_identified_per_row(db, graded):
    from app.features.assessments.models import GradePublication
    full_term(graded)
    publish_grades(graded, "midterm")
    s2 = graded["students"][1]
    first = db.scalar(select(GradePublication))
    older = {**first.snapshot["policy"], "passing": 60, "categories": [
        {"key": "participation", "label": "Participation", "weight": 100}]}
    db.add(GradePublication(offering_id=first.offering_id, student_id=s2.id, period="midterm", release_number=1,
                            grade=70, remark="passed", policy_version=7, published_by=first.published_by,
                            snapshot={"policy_version": 7, "policy": older, "periods": {
                                "midterm": {"grade": "70", "categories": [{"key": "participation", "percent": "70.00"}]}}}))
    db.commit()
    _, meta, header, rows = sheet_rows(export(graded, basis="published"))
    assert "Policy version" in header and "Participation (%)" in header and "Quizzes (%)" in header
    by_name = {r[1]: dict(zip(header, r, strict=True)) for r in rows}
    assert by_name["s1"]["Policy version"] == 1 and by_name["s2"]["Policy version"] == 7
    assert by_name["s2"]["Participation (%)"] == 70.0 and by_name["s2"]["Quizzes (%)"] is None
    assert "version 1;" in meta["Grading policy"] and "version 7; passing grade 60" in meta["Grading policy"]


def test_headers_and_labels_are_formula_protected_too(graded):
    from helpers import POLICY, publish_syllabus
    full_term(graded)
    labelled = {**POLICY, "categories": [{**c, "label": '=cmd|"/c calc"!A1' if c["key"] == "quiz" else c["label"]}
                                         for c in POLICY["categories"]]}
    publish_syllabus(graded, policy=labelled)
    ws = load_workbook(BytesIO(export(graded, basis="working").content)).active
    assert not any(c.data_type == "f" for row in ws.iter_rows() for c in row)


@pytest.mark.parametrize("columns", [8, 9, 11, 12, 14, 16])
def test_pdf_columns_always_fit_the_page_and_names_stay_wide(columns):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    widths = render.column_widths(columns)
    assert len(widths) == columns and abs(sum(widths) - (landscape(A4)[0] - 20 * mm)) < 1
    assert widths[1] >= 32 * mm                        # a long name never has to break inside a word
    assert min(widths) >= 12 * mm                      # and no column collapses


def test_a_course_export_reproduces_its_own_release_and_period_republishing_keeps_its_review_flag(graded):
    ids = full_term(graded)
    for period in ("midterm", "finals", "course"):
        publish_grades(graded, period)
    assert score(graded, ids["exam_m"], graded["students"][0], 44).status_code == 200     # midterm 81 -> 85, course 87 -> 89
    publish_grades(graded, "midterm")                                                       # the midterm release moves to 85
    rows, _ = book(graded)
    assert rows["s1"]["published"]["midterm"]["needs_review"] is False
    assert rows["s1"]["published"]["course"]["needs_review"] is True                        # only a Course republication clears it
    _, _, header, data = sheet_rows(export(graded, period="course", basis="published"))
    s1 = dict(zip(header, next(r for r in data if r[1] == "s1"), strict=True))
    assert s1["Midterm grade (used for this Course release)"] == 81.0                     # what the Course grade was built from,
    assert s1["Grade"] == 87.0 and "review pending" in s1["Status"]                       # not the newer 85 midterm release
    _, _, mid_header, mid = sheet_rows(export(graded, period="midterm", basis="published"))
    assert dict(zip(mid_header, next(r for r in mid if r[1] == "s1"), strict=True))["Grade"] == 85.0
    publish_grades(graded, "course")
    _, _, header, data = sheet_rows(export(graded, period="course", basis="published"))
    s1 = dict(zip(header, next(r for r in data if r[1] == "s1"), strict=True))
    assert s1["Midterm grade (used for this Course release)"] == 85.0 and s1["Grade"] == 89.0 and "review" not in s1["Status"]


def test_a_students_standing_pdf_has_a_published_and_a_faculty_basis(graded):
    ids = full_term(graded)
    publish_grades(graded, "midterm")
    assert score(graded, ids["exam_m"], graded["students"][0], 44).status_code == 200      # working 85 vs released 81
    sid = graded["students"][0].id
    url = lambda basis: g(graded, f"/students/{sid}/standing.pdf?basis={basis}")
    handout = graded["fac"].get(url("published"))
    assert handout.status_code == 200 and handout.headers["content-type"] == "application/pdf"
    assert handout.content[:5] == b"%PDF-" and "FACULTY-COPY" not in handout.headers["content-disposition"]
    text = " ".join(pdf_pages(handout))
    assert "FACULTY COPY" not in text and "Working score" not in text and "Current calculation" not in text
    assert "81.00" in text and "Visible to student" in text                                 # only the release
    faculty_copy = graded["fac"].get(url("working"))
    pages = pdf_pages(faculty_copy)
    assert "FACULTY-COPY" in faculty_copy.headers["content-disposition"] and all("FACULTY COPY" in p for p in pages)
    assert "Working score" in " ".join(pages) and "85.00" in " ".join(pages)
    # the same people as the page itself, and the export is audited without any grades in the log
    assert graded["st1"].get(url("published")).status_code == 403 and graded["admin"].get(url("published")).status_code == 403
    assert graded["fac"].get(g(graded, f"/students/{uuid.uuid4()}/standing.pdf")).status_code == 404
    events = [e for e in graded["admin"].get("/api/audit").json() if e["action"] == "standing.exported"]
    assert len(events) == 2 and all(set(e["details"]) == {"student_id", "basis"} for e in events) and {e["details"]["basis"] for e in events} == {"published", "working"}


def test_class_standing_ranks_with_ties_keeps_unranked_apart_and_matches_the_export(db, graded):
    full_term(graded)
    publish_grades(graded, "midterm")                                    # s1 released, s2 not
    admin = graded["admin"]
    s3 = make_account(db, "s3@example.com", "student", "S-3")
    s4 = make_account(db, "s4@example.com", "student", "S-4")
    for acct in (s3, s4):
        admin.post(f"/api/sections/{graded['sec_a']['id']}/members", {"student_id": str(acct.id)})
    url = lambda **q: g(graded, "/class-standing?" + "&".join(f"{k}={v}" for k, v in {"period": "midterm", **q}.items()))
    page = graded["fac"].get(url()).json()
    assert page["columns"][0] == "Rank" and page["banner"] is None                      # published basis is not a preview
    by_name = {r[2]: r for r in page["rows"]}
    assert by_name["s1"][0] == 1 and by_name["s2"][0] is None and by_name["s3"][0] is None   # unpublished is unranked, not last
    assert page["ranking"]["ranked"] == 1 and page["ranking"]["unranked"] == 3 and page["ranking"]["scope"] == "all sections of this offering"
    assert next(r[0] for r in page["rows"]) == 1 and len(page["student_ids"]) == len(page["rows"])
    # a section scope re-ranks inside that section only
    one = graded["fac"].get(url(section_id=graded["sec_b"]["id"])).json()
    assert [r[2] for r in one["rows"]] == ["s2"] and one["ranking"]["scope"] == "1B"
    # working basis: everyone with a complete calculation is ranked; ties share a rank (1, 2, 2, 4)
    from app.features.exports.tables import competition_ranks
    assert competition_ranks([90, 85, 85, 70, None, 0]) == [1, 2, 2, 4, None, 5]          # an explicit zero is ranked
    preview = graded["fac"].get(url(basis="working")).json()
    assert preview["banner"] and preview["ranking"]["ranked"] + preview["ranking"]["unranked"] == 4
    # the export agrees with the page (same rows, rank first) and says how the rank was made
    sheet = export(graded, basis="published", rank="true")
    _, meta, header, rows = sheet_rows(sheet)
    assert header[0] == "Rank" and [r[0] for r in rows] == [r[0] for r in page["rows"]]
    assert "Rank within all sections" in meta["Ranking"] and "competition ranking" in meta["Ranking"]
    pdf = pdf_pages(export(graded, basis="working", format="pdf", rank="true"))
    assert all("PREVIEW" in p for p in pdf) and "Rank" in pdf[0]
    assert "Rank" not in sheet_rows(export(graded, basis="published"))[2][0]                 # plain export is unchanged
    # faculty only: students and administrators get nothing, an unrelated teacher gets nothing
    assert graded["st1"].get(url()).status_code == 403 and admin.get(url()).status_code == 403
    stranger = Api(make_account(db, "rank-stranger@example.com", "faculty").email)
    assert stranger.get(url()).status_code in (403, 404)

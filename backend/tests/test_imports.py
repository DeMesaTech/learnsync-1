import io

import openpyxl
import pytest
from docx import Document
from helpers import Api, make_account
from sqlalchemy import func, select

from app.config import settings
from app.features.academics import documents, imports
from app.features.academics.models import Subject
from app.features.accounts import faculty_imports
from app.features.accounts.models import Account


@pytest.fixture
def stage(db, actors, tmp_path, monkeypatch):
    """Isolated upload dir, captured invitation emails, an admin, and a term with sections."""
    monkeypatch.setattr(settings(), "upload_root", tmp_path / "uploads")
    sent = []
    monkeypatch.setattr(imports, "send_link", lambda account, token, purpose: sent.append(account.email))
    monkeypatch.setattr(faculty_imports, "send_link", lambda account, token, purpose: sent.append(account.email))
    admin = Api("admin@example.com")
    college = admin.post("/api/colleges", {
        "code": "BSE", "name": "College of Business and Entrepreneurship"}).json()
    year = admin.post("/api/school-years", {
        "label": "2026-2027", "start_date": "2026-06-01", "end_date": "2027-04-30",
        "terms": [{"name": "First Semester", "sequence": 1, "start_date": "2026-06-01",
                   "end_date": "2026-10-31"}]}).json()
    term = year["terms"][0]
    sections = {n: admin.post(f"/api/terms/{term['id']}/sections",
                              {"name": n, "year_level": 1}).json() for n in ("1A", "1B")}
    return {"admin": admin, "college": college, "term": term, "sections": sections, "sent": sent,
            "root": tmp_path / "uploads"}


CSV = (b"student_number,email,display_name,section\r\n"
       b"2026-001,ana@example.com,Ana Reyes,1A\r\n"
       b"2026-002,ben@example.com,Ben Cruz,1A\r\n"
       b"2026-003,cara@example.com,Cara Dela Cruz,\r\n")


def upload(stage, content=CSV, name="students.csv", term=True):
    data = {"term_id": stage["term"]["id"]} if term else {}
    return stage["admin"].upload("/api/student-imports", name, content, data)


def account_count(db):
    return db.scalar(select(func.count()).select_from(Account))


def prospectus_commit(stage, expected_revision, college_id=None):
    return {"expected_revision": expected_revision,
            "college_id": college_id or stage["college"]["id"]}


def test_student_import_preview_then_commit(db, stage):
    before = account_count(db)
    response = upload(stage)
    assert response.status_code == 201, response.text
    preview = response.json()
    assert preview["summary"] == {"new": 3, "existing": 0, "errors": 0, "total": 3}
    assert account_count(db) == before            # preview creates nothing
    done = stage["admin"].post(f"/api/student-imports/{preview['id']}/commit")
    assert done.status_code == 200, done.text
    assert done.json() == {"created": 3, "existing": 0, "skipped": 0, "email_failed": []}
    assert sorted(stage["sent"]) == ["ana@example.com", "ben@example.com", "cara@example.com"]
    ana = db.scalar(select(Account).where(Account.email == "ana@example.com"))
    assert ana.role == "student" and ana.status == "invited" and ana.student_number == "2026-001"
    members = stage["admin"].get(f"/api/sections/{stage['sections']['1A']['id']}/members").json()
    assert sorted(m["display_name"] for m in members) == ["Ana Reyes", "Ben Cruz"]
    # committing twice is refused
    assert stage["admin"].post(f"/api/student-imports/{preview['id']}/commit").status_code == 409


def test_student_import_skips_error_rows_and_commits_valid_rows(db, stage):
    make_account(db, "taken@example.com", "faculty")
    bad = (b"student_number,email,display_name,section\n"
           b"2026-001,ok@example.com,Fine Student,1A\n"
           b"2026-002,not-an-email,Bad Email,1A\n"
           b"2026-003,ok@example.com,Repeated Email,1A\n"
           b"2026-004,taken@example.com,Belongs To Faculty,1A\n"
           b"2026-005,five@example.com,Unknown Section,9Z\n")
    before = account_count(db)
    preview = upload(stage, bad).json()
    assert preview["summary"]["errors"] == 4 and preview["summary"]["new"] == 1
    by_row = {r["row"]: r for r in preview["rows"]}
    assert by_row[2]["status"] == "new"
    assert "not valid" in by_row[3]["errors"][0]
    assert "repeats row 2" in by_row[4]["errors"][0]
    assert "different account" in by_row[5]["errors"][0]
    assert "does not exist" in by_row[6]["errors"][0]
    done = stage["admin"].post(f"/api/student-imports/{preview['id']}/commit")
    assert done.status_code == 200, done.text
    assert done.json() == {"created": 1, "existing": 0, "skipped": 4, "email_failed": []}
    assert account_count(db) == before + 1
    assert stage["sent"] == ["ok@example.com"]
    member = db.scalar(select(Account).where(Account.email == "ok@example.com"))
    assert member.display_name == "Fine Student"
    assert stage["admin"].post(f"/api/student-imports/{preview['id']}/commit").status_code == 409


def test_student_import_refuses_when_every_row_has_errors(db, stage):
    content = (b"student_number,email,display_name,section\n"
               b"2026-010,bad-email,Invalid Student,9Z\n")
    preview = upload(stage, content).json()
    before = account_count(db)

    refused = stage["admin"].post(f"/api/student-imports/{preview['id']}/commit")

    assert refused.status_code == 409
    assert refused.json()["error"]["code"] == "import_no_valid_rows"
    assert account_count(db) == before and stage["sent"] == []


def test_existing_student_is_reused_and_section_conflict_is_flagged(db, stage):
    first = upload(stage).json()
    stage["admin"].post(f"/api/student-imports/{first['id']}/commit")
    again = (b"student_number,email,display_name,section\n"
             b"2026-001,ana@example.com,Ana Reyes,1A\n"        # already in 1A -> fine
             b"2026-002,ben@example.com,Ben Cruz,1B\n")        # already in 1A -> conflict
    preview = upload(stage, again).json()
    assert [r["status"] for r in preview["rows"]] == ["existing", "error"]
    assert "different section" in preview["rows"][1]["errors"][0]


def test_reimporting_the_same_file_is_a_no_op(db, stage):
    first = upload(stage).json()
    stage["admin"].post(f"/api/student-imports/{first['id']}/commit")
    stage["sent"].clear()
    again = upload(stage).json()
    assert again["summary"] == {"new": 0, "existing": 3, "errors": 0, "total": 3}
    done = stage["admin"].post(f"/api/student-imports/{again['id']}/commit")
    assert done.status_code == 200, done.text
    assert done.json() == {"created": 0, "existing": 3, "skipped": 0, "email_failed": []}
    assert stage["sent"] == []                                  # nobody is re-invited
    members = stage["admin"].get(f"/api/sections/{stage['sections']['1A']['id']}/members").json()
    assert len(members) == 2                                    # no duplicate placement


def test_xlsx_import_and_header_aliases(db, stage):
    book = openpyxl.Workbook()
    book.active.append(["Student No", "Email Address", "Name"])
    book.active.append(["2026-100", "xl@example.com", "Xl Student"])
    buffer = io.BytesIO()
    book.save(buffer)
    preview = upload(stage, buffer.getvalue(), "list.xlsx", term=False).json()
    assert preview["summary"]["new"] == 1


@pytest.mark.parametrize("middle_header", ["MiddleInitial", "Middle Initial"])
def test_student_import_combines_split_name_columns(tmp_path, middle_header):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Student Number", "email", "lastName", "firstName", middle_header])
    sheet.append(["2026-200", "juan@example.com", "Dela Cruz", "Juan", "Adan"])
    sheet.append(["2026-201", "ana@example.com", "Reyes,", "Ana", None])
    path = tmp_path / "split-names.xlsx"
    book.save(path)

    rows = documents.read_student_rows(path, ".xlsx")

    assert [row["display_name"] for row in rows] == ["Juan Adan Dela Cruz", "Ana Reyes"]


def test_student_xlsx_import_reads_each_sheet_and_keeps_source_location(tmp_path):
    book = openpyxl.Workbook()
    first = book.active
    first.title = "1A"
    first.append(["Student Number", "email", "lastName", "firstName", "MiddleInitial"])
    first.append(["S-1", "one@example.com", "Cruz", "Ana", "Reyes"])
    second = book.create_sheet("2A")
    second.append(["Section 2A roster"])
    second.append([])
    second.append(["Student Number", "email", "lastName", "firstName", "Middle Initial"])
    second.append(["S-2", "two@example.com", "Santos", "Ben", None])
    path = tmp_path / "sectioned-students.xlsx"
    book.save(path)

    rows = documents.read_student_rows(path, ".xlsx")

    assert [row["display_name"] for row in rows] == ["Ana Reyes Cruz", "Ben Santos"]
    assert [(row["source_sheet"], row["source_row"]) for row in rows] == [
        ("1A", 2), ("2A", 4)]


def test_file_validation(db, stage):
    assert upload(stage, CSV, "students.exe").status_code == 422
    assert upload(stage, b"this is not a zip", "students.xlsx").status_code == 422
    assert upload(stage, b"", "students.csv").status_code == 422
    missing = upload(stage, b"email,name\nx@example.com,X\n")
    assert missing.status_code == 422 and missing.json()["error"]["code"] == "columns_missing"
    assert upload(stage, "é,\xff".encode("latin-1") + b"\n", "s.csv").status_code == 422
    # sections need a term
    no_term = upload(stage, CSV, term=False).json()
    assert no_term["summary"]["errors"] == 2
    # nothing failed is left behind on disk
    leftovers = [p for p in stage["root"].rglob("*") if p.is_file() and p.suffix == ".part"]
    assert leftovers == []


def test_uploads_are_private_and_admin_only(db, stage):
    preview = upload(stage).json()
    assert preview["id"]
    stored = [p for p in stage["root"].rglob("*") if p.is_file()]
    assert len(stored) == 1 and "students" not in stored[0].name   # generated key, not user name
    faculty = Api("faculty@example.com")
    assert faculty.upload("/api/student-imports", "s.csv", CSV).status_code == 403
    assert faculty.get(f"/api/student-imports/{preview['id']}").status_code == 403


FACULTY_CSV = (b"email,display_name\r\n"
               b"teacher.one@example.com,Teacher One\r\n"
               b"teacher.two@example.com,Teacher Two\r\n")


def faculty_upload(stage, content=FACULTY_CSV, name="teachers.csv"):
    return stage["admin"].upload("/api/faculty-imports", name, content)


def test_faculty_import_preview_and_commit(db, stage):
    before = account_count(db)
    response = faculty_upload(stage)
    assert response.status_code == 201, response.text
    preview = response.json()
    assert preview["summary"] == {"new": 2, "existing": 0, "errors": 0, "total": 2}
    assert account_count(db) == before
    done = stage["admin"].post(f"/api/faculty-imports/{preview['id']}/commit")
    assert done.status_code == 200, done.text
    assert done.json() == {"created": 2, "existing": 0, "email_failed": []}
    assert sorted(stage["sent"]) == ["teacher.one@example.com", "teacher.two@example.com"]
    teacher = db.scalar(select(Account).where(Account.email == "teacher.one@example.com"))
    assert teacher.role == "faculty" and teacher.status == "invited"
    assert stage["admin"].post(f"/api/faculty-imports/{preview['id']}/commit").status_code == 409


def test_faculty_import_reports_duplicates_and_non_faculty_accounts_without_partial_create(db, stage):
    make_account(db, "taken@example.com", "student", student_number="S-1")
    before = account_count(db)
    content = (b"email,display_name\n"
               b"first@example.com,First Teacher\n"
               b"not-an-email,Invalid Teacher\n"
               b"first@example.com,Duplicate Teacher\n"
               b"taken@example.com,Existing Student\n")
    preview = faculty_upload(stage, content).json()
    assert preview["summary"] == {"new": 1, "existing": 0, "errors": 3, "total": 4}
    rows = {row["row"]: row for row in preview["rows"]}
    assert "not valid" in rows[3]["errors"][0]
    assert "repeats row 2" in rows[4]["errors"][0]
    assert "different account" in rows[5]["errors"][0]
    response = stage["admin"].post(f"/api/faculty-imports/{preview['id']}/commit")
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "import_has_errors"
    assert account_count(db) == before and stage["sent"] == []


def test_faculty_import_preview_edits_are_revalidated_and_persisted(stage):
    preview = faculty_upload(stage, b"email,display_name\nbroken-email,New Teacher\n").json()
    saved = stage["admin"].put(f"/api/faculty-imports/{preview['id']}", {"rows": [
        {"row": 2, "display_name": "Corrected Teacher", "email": "corrected@example.com"},
    ]})
    assert saved.status_code == 200, saved.text
    assert saved.json()["summary"] == {"new": 1, "existing": 0, "errors": 0, "total": 1}
    assert saved.json()["rows"][0]["display_name"] == "Corrected Teacher"
    assert stage["sent"] == []

    done = stage["admin"].post(f"/api/faculty-imports/{preview['id']}/commit")
    assert done.status_code == 200, done.text
    assert stage["sent"] == ["corrected@example.com"]
    assert stage["admin"].put(f"/api/faculty-imports/{preview['id']}", {"rows": [
        {"row": 2, "display_name": "Changed after commit", "email": "corrected@example.com"},
    ]}).status_code == 409


def test_faculty_import_skips_existing_faculty_and_supports_xlsx(db, stage):
    make_account(db, "existing-teacher@example.com", "faculty")
    book = openpyxl.Workbook()
    book.active.append(["Email Address", "Teacher Name"])
    book.active.append(["existing-teacher@example.com", "Existing Teacher"])
    book.active.append(["new-teacher@example.com", "New Teacher"])
    buffer = io.BytesIO()
    book.save(buffer)
    response = faculty_upload(stage, buffer.getvalue(), "teachers.xlsx")
    assert response.status_code == 201, response.text
    preview = response.json()
    assert preview["summary"] == {"new": 1, "existing": 1, "errors": 0, "total": 2}
    done = stage["admin"].post(f"/api/faculty-imports/{preview['id']}/commit")
    assert done.status_code == 200
    assert done.json() == {"created": 1, "existing": 1, "email_failed": []}
    assert stage["sent"] == ["new-teacher@example.com"]


def test_faculty_import_validation_and_role_access(stage):
    assert faculty_upload(stage, FACULTY_CSV, "teachers.exe").status_code == 422
    assert faculty_upload(stage, b"email,name\nx@example.com,X\n").status_code == 422
    faculty = Api("faculty@example.com")
    assert faculty.upload("/api/faculty-imports", "teachers.csv", FACULTY_CSV).status_code == 403


# ---------- prospectus ----------

def test_docx_subject_rows_keep_their_year_and_semester_heading(tmp_path):
    doc = Document()
    sections = [
        ("1st YEAR — 1st Semester", [("BCA 1", "Accounting Plus", "6"),
                                    ("GEN ED 001", "Purposive Communication", "3")]),
        ("1st YEAR — 2nd Semester", [("AE 113", "Financial Accounting & Reporting", "6")]),
        ("2nd YEAR — 1st Semester", [("EnCC 213", "Market Research and Consumer", "3")]),
        ("2nd YEAR — 2nd Semester", [("EnCC 214", "Opportunity Seeking", "3")]),
        ("3rd YEAR — 1st Semester", [("EnCC 314", "Business Plan Preparation I", "4")]),
        ("3rd YEAR — 2nd Semester", [("SPT ELEC 343", "Franchising and Retail Management", "3")]),
    ]
    for heading, subjects in sections:
        doc.add_paragraph(heading)
        table = doc.add_table(rows=1, cols=5)
        for cell, value in zip(table.rows[0].cells,
                               ("Subject Code", "Descriptive Title", "Lec", "Lab", "Units")):
            cell.text = value
        for code, title, units in subjects:
            for cell, value in zip(table.add_row().cells, (code, title, "3", "0", units)):
                cell.text = value

    path = tmp_path / "prospectus.docx"
    doc.save(path)
    draft, warnings = documents.parse_prospectus(documents.read_lines(path, ".docx"))

    assert [(row["code"], row["year_level"], row["semester"]) for row in draft] == [
        ("BCA 1", 1, 1),
        ("GEN ED 001", 1, 1),
        ("AE 113", 1, 2),
        ("ENCC 213", 2, 1),
        ("ENCC 214", 2, 2),
        ("ENCC 314", 3, 1),
        ("SPT ELEC 343", 3, 2),
    ]
    assert warnings == []


def test_docx_table_without_a_clear_year_and_semester_heading_is_unplaced(tmp_path):
    doc = Document()
    doc.add_paragraph("1st YEAR — 1st Semester")
    for code, title in (("ENT 101", "Recognized Subject"),):
        table = doc.add_table(rows=1, cols=3)
        for cell, value in zip(table.rows[0].cells,
                               ("Subject Code", "Descriptive Title", "Units")):
            cell.text = value
        for cell, value in zip(table.add_row().cells, (code, title, "3")):
            cell.text = value
    table = doc.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells,
                           ("Subject Code", "Descriptive Title", "Units")):
        cell.text = value
    for cell, value in zip(table.add_row().cells, ("ENT 102", "Unplaced Subject", "4")):
        cell.text = value

    path = tmp_path / "missing-heading.docx"
    doc.save(path)
    draft, warnings = documents.parse_prospectus(documents.read_lines(path, ".docx"))

    assert [(row["code"], row["year_level"], row["semester"]) for row in draft] == [
        ("ENT 101", 1, 1), ("ENT 102", None, None)]
    assert any("no year level/semester" in warning for warning in warnings)


def docx_bytes(lines):
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


PROSPECTUS = ["BS Entrepreneurship Prospectus", "First Year", "First Semester",
              "ENT 101 Introduction to Entrepreneurship 3 3 0",
              "GE 101 Understanding the Self 3", "Second Semester",
              "ENT 102 Business Opportunity Analysis 3", "something odd ENT 9 broken"]


def test_prospectus_extraction_review_and_commit(db, stage):
    admin = stage["admin"]
    db.add(Subject(code="GE 101", title="Existing", units=3))
    db.commit()
    created = admin.upload("/api/prospectus-imports", "prospectus.docx", docx_bytes(PROSPECTUS))
    assert created.status_code == 201, created.text
    draft = created.json()
    rows = {r["code"]: r for r in draft["subjects"]}
    assert rows["ENT 101"]["title"] == "Introduction to Entrepreneurship"
    assert rows["ENT 101"]["units"] == "3" and rows["ENT 101"]["year_level"] == 1
    assert rows["ENT 102"]["semester"] == 2 and rows["GE 101"]["exists"] is True
    assert draft["revision"] == 1
    assert db.scalar(select(func.count()).select_from(Subject)) == 1     # nothing created yet

    edited = [{"code": r["code"], "title": r["title"], "units": r["units"],
               "year_level": r["year_level"], "semester": r["semester"]}
              for r in draft["subjects"]]
    edited[0]["title"] = "Intro to Entrepreneurship (reviewed)"
    saved = admin.put(f"/api/prospectus-imports/{draft['id']}",
                      {"expected_revision": 1, "subjects": edited})
    assert saved.status_code == 200 and saved.json()["revision"] == 2
    stale = admin.put(f"/api/prospectus-imports/{draft['id']}",
                      {"expected_revision": 1, "subjects": edited})
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "draft_revision_conflict"

    done = admin.post(f"/api/prospectus-imports/{draft['id']}/commit",
                      prospectus_commit(stage, 2))
    assert done.status_code == 200 and done.json() == {"created": 2, "skipped_existing": 1}
    subject = db.scalar(select(Subject).where(Subject.code == "ENT 101"))
    assert subject.title == "Intro to Entrepreneurship (reviewed)"
    assert str(subject.college_id) == stage["college"]["id"]
    existing = db.scalar(select(Subject).where(Subject.code == "GE 101"))
    assert existing.title == "Existing"                                   # never overwritten
    assert existing.college_id is None
    again = admin.post(f"/api/prospectus-imports/{draft['id']}/commit",
                       prospectus_commit(stage, 2))
    assert again.status_code == 409


def test_prospectus_commit_requires_an_active_college(db, stage):
    admin = stage["admin"]
    draft = admin.upload("/api/prospectus-imports", "prospectus.docx",
                         docx_bytes(PROSPECTUS)).json()
    inactive = admin.post("/api/colleges", {
        "code": "OLD", "name": "Inactive College", "status": "inactive"}).json()

    refused = admin.post(f"/api/prospectus-imports/{draft['id']}/commit",
                         prospectus_commit(stage, draft["revision"], inactive["id"]))
    assert refused.status_code == 409
    assert refused.json()["error"]["code"] == "college_inactive"
    assert db.scalar(select(func.count()).select_from(Subject)) == 0

    committed = admin.post(f"/api/prospectus-imports/{draft['id']}/commit",
                           prospectus_commit(stage, draft["revision"]))
    assert committed.status_code == 200
    assert all(str(subject.college_id) == stage["college"]["id"]
               for subject in db.scalars(select(Subject)).all())


def test_invalid_draft_cannot_commit_and_manual_entry_works(db, stage):
    admin = stage["admin"]
    manual = admin.post("/api/prospectus-imports")
    assert manual.status_code == 201 and manual.json()["subjects"] == []
    pid = manual.json()["id"]
    assert admin.post(f"/api/prospectus-imports/{pid}/commit",
                      prospectus_commit(stage, 1)).status_code == 422
    admin.put(f"/api/prospectus-imports/{pid}", {"expected_revision": 1, "subjects": [
        {"code": "ENT 201", "title": "Marketing", "units": "3", "year_level": 2, "semester": 1},
        {"code": "", "title": "No code", "units": "abc"}]})
    refused = admin.post(f"/api/prospectus-imports/{pid}/commit",
                         prospectus_commit(stage, 2))
    assert refused.status_code == 409 and refused.json()["error"]["code"] == "draft_has_errors"
    assert db.scalar(select(func.count()).select_from(Subject)) == 0     # no partial commit


def test_unreadable_prospectus_falls_back_to_manual(db, stage):
    blank = stage["admin"].upload("/api/prospectus-imports", "scan.docx", docx_bytes(["Nothing"]))
    assert blank.status_code == 201
    assert blank.json()["subjects"] == [] and any("manually" in w for w in blank.json()["warnings"])
    fake_pdf = stage["admin"].upload("/api/prospectus-imports", "scan.pdf", b"not a pdf at all")
    assert fake_pdf.status_code == 422

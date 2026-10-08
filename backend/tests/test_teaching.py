import io
import uuid

from docx import Document
from fastapi.testclient import TestClient
from helpers import PDF, Api, learn, make_account, make_outline, publish_syllabus, t, uid

from app.main import app


def new_item(c, kind="lesson", title="Lesson 1", **extra):
    response = c["fac"].post(t(c, "/items"), {"kind": kind, "title": title, **extra})
    assert response.status_code == 201, response.text
    return response.json()


def edit_and_publish(c, item_id, **fields):
    draft = c["fac"].get(t(c, f"/items/{item_id}")).json()["draft"]
    saved = c["fac"].put(t(c, f"/items/{item_id}/draft"),
                         {"expected_counter": draft["counter"], **fields})
    assert saved.status_code == 200, saved.text
    done = c["fac"].post(t(c, f"/items/{item_id}/draft/publish"),
                         {"expected_counter": saved.json()["counter"]})
    return done


# ---------------- access ----------------

def test_only_the_assigned_faculty_can_teach(db, course):
    other = make_account(db, "other@example.com", "faculty")
    stranger = Api(other.email)
    assert stranger.get(t(course, "/syllabus")).status_code == 404
    assert stranger.post(t(course, "/items"), {"kind": "lesson", "title": "x"}).status_code == 404
    assert course["st1"].get(t(course, "/syllabus")).status_code == 403     # students: no teach API
    assert course["admin"].get(t(course, "/syllabus")).status_code == 403   # admin: no teaching
    assert TestClient(app).get(t(course, "/syllabus")).status_code == 401
    assert course["fac"].get(learn(course, "/syllabus")).status_code == 403  # faculty: no learn API


# ---------------- syllabus ----------------

def test_syllabus_draft_publish_and_student_visibility(course):
    fac, st1 = course["fac"], course["st1"]
    assert st1.get(learn(course, "/syllabus")).json() == {"published": None, "covered": []}
    draft = fac.post(t(course, "/syllabus/draft")).json()
    assert draft["outline"]["course"]["name"] == "Intro to Entrepreneurship"   # prefilled from subject
    assert draft["counter"] == 1
    assert fac.post(t(course, "/syllabus/draft")).json()["id"] == draft["id"]   # idempotent
    # a draft is invisible to students
    assert st1.get(learn(course, "/syllabus")).json()["published"] is None
    # incomplete outline cannot be published
    refused = fac.post(t(course, "/syllabus/draft/publish"), {"expected_counter": 1})
    assert refused.status_code == 422 and refused.json()["error"]["code"] == "syllabus_incomplete"
    outline, _, topic_id = make_outline()
    saved = fac.put(t(course, "/syllabus/draft"), {"expected_counter": 1, "outline": outline})
    assert saved.json()["counter"] == 2
    assert fac.post(t(course, "/syllabus/draft/publish"),
                    {"expected_counter": 2}).json()["version"] == 1
    published = st1.get(learn(course, "/syllabus")).json()["published"]
    assert published["version"] == 1 and published["outline"]["chapters"][0]["topics"][0]["id"] == topic_id
    # editing creates a new draft; students keep seeing the published revision meanwhile
    draft2 = fac.post(t(course, "/syllabus/draft")).json()
    assert draft2["version"] == 2 and draft2["outline"]["chapters"][0]["title"] == outline["chapters"][0]["title"]
    outline2 = {**outline, "course": {**outline["course"], "name": "Renamed course"}}
    fac.put(t(course, "/syllabus/draft"), {"expected_counter": 1, "outline": outline2})
    assert st1.get(learn(course, "/syllabus")).json()["published"]["outline"]["course"]["name"] \
        == "Intro to Entrepreneurship"
    fac.post(t(course, "/syllabus/draft/publish"), {"expected_counter": 2})
    state = fac.get(t(course, "/syllabus")).json()
    assert state["published"]["version"] == 2 and state["draft"] is None
    assert [(h["version"], h["state"]) for h in state["history"]] == [(2, "published"), (1, "superseded")]
    assert st1.get(learn(course, "/syllabus")).json()["published"]["version"] == 2


def test_second_tab_cannot_silently_overwrite(course):
    fac_a, fac_b = course["fac"], Api("faculty@example.com")     # two browser tabs
    draft = fac_a.post(t(course, "/syllabus/draft")).json()
    outline, _, _ = make_outline()
    ok = fac_a.put(t(course, "/syllabus/draft"), {"expected_counter": draft["counter"], "outline": outline})
    assert ok.status_code == 200
    stale = fac_b.put(t(course, "/syllabus/draft"), {"expected_counter": draft["counter"], "outline": outline})
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "draft_revision_conflict"
    # the newer server draft is still readable for reconciliation
    assert fac_b.get(t(course, "/syllabus")).json()["draft"]["counter"] == ok.json()["counter"]
    assert fac_b.delete(t(course, f"/syllabus/draft?expected_counter={draft['counter']}")).status_code == 409


def test_discard_draft_keeps_published(course):
    publish_syllabus(course)
    fac = course["fac"]
    draft = fac.post(t(course, "/syllabus/draft")).json()
    assert fac.delete(t(course, f"/syllabus/draft?expected_counter={draft['counter']}")).status_code == 204
    state = fac.get(t(course, "/syllabus")).json()
    assert state["draft"] is None and state["published"]["version"] == 1


def test_invalid_grading_policy_blocks_publish_but_valid_one_passes(course):
    fac = course["fac"]
    outline, _, _ = make_outline()
    draft = fac.post(t(course, "/syllabus/draft")).json()
    bad = {"categories": [{"key": "quiz", "label": "Quizzes", "weight": 60},
                          {"key": "exam", "label": "Exams", "weight": 30}]}
    saved = fac.put(t(course, "/syllabus/draft"), {"expected_counter": draft["counter"],
                                                    "outline": outline, "grading_policy": bad}).json()
    refused = fac.post(t(course, "/syllabus/draft/publish"), {"expected_counter": saved["counter"]})
    assert refused.status_code == 422 and "100%" in refused.json()["error"]["message"]
    good = {"categories": [{"key": "quiz", "label": "Quizzes", "weight": 50},
                           {"key": "exam", "label": "Exams", "weight": 50}],
            "periods": [{"key": "midterm", "label": "Midterm", "share": 50},
                        {"key": "finals", "label": "Finals", "share": 50}]}
    saved = fac.put(t(course, "/syllabus/draft"), {"expected_counter": saved["counter"],
                                                    "outline": outline, "grading_policy": good}).json()
    assert fac.post(t(course, "/syllabus/draft/publish"),
                    {"expected_counter": saved["counter"]}).status_code == 200


def test_outline_validation_rejects_duplicate_ids(course):
    fac = course["fac"]
    outline, chapter_id, _ = make_outline()
    outline["chapters"].append({"id": chapter_id, "title": "Clone"})
    draft = fac.post(t(course, "/syllabus/draft")).json()
    response = fac.put(t(course, "/syllabus/draft"), {"expected_counter": draft["counter"],
                                                       "outline": outline})
    assert response.status_code == 422


# ---------------- syllabus import ----------------

def docx_syllabus():
    doc = Document()
    info = doc.add_table(rows=3, cols=2)
    for row, (k, v) in zip(info.rows, [("1. Course No.", "ENT 101"),
                                       ("2. Course Name", "Intro to Entrepreneurship"),
                                       ("4. Credit Units", "3")], strict=True):
        row.cells[0].text, row.cells[1].text = k, v
    cov = doc.add_table(rows=3, cols=5)
    rows = [["Week 1", "Understand the basics", "CHAPTER 1 - Foundations\nWhat is a venture?\nOpportunity",
             "Lecture\nDiscussion", "Quiz 1"],
            ["Week 2 - 4", "Apply tools", "CHAPTER 2 - Planning\nBusiness model", "Case study", "Plan"],
            ["Week 5", "MIDTERM EXAMINATION", "", "", ""]]
    for row, values in zip(cov.rows, rows, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_syllabus_import_reads_a_template_into_a_reviewable_draft(course):
    fac = course["fac"]
    response = fac.upload(t(course, "/syllabus/import"), "syllabus.docx", docx_syllabus())
    assert response.status_code == 201, response.text
    draft = response.json()["draft"]
    chapters = draft["outline"]["chapters"]
    assert draft["outline"]["course"]["code"] == "ENT 101" and draft["outline"]["course"]["units"] == "3"
    assert [c["title"] for c in chapters if c["kind"] == "chapter"] == \
        ["CHAPTER 1 - Foundations", "CHAPTER 2 - Planning"]
    assert chapters[0]["weeks"] == "Week 1" and chapters[0]["assessment"] == "Quiz 1"
    assert [x["title"] for x in chapters[0]["topics"]] == ["What is a venture?", "Opportunity"]
    assert chapters[2]["kind"] == "exam"
    assert draft["source_file"]["name"] == "syllabus.docx"
    # the draft is not published until faculty reviews it; a second import needs the draft resolved
    assert course["st1"].get(learn(course, "/syllabus")).json()["published"] is None
    again = fac.upload(t(course, "/syllabus/import"), "syllabus.docx", docx_syllabus())
    assert again.status_code == 409 and again.json()["error"]["code"] == "draft_exists"


def test_unreadable_syllabus_falls_back_to_manual_editing(course):
    doc = Document()
    doc.add_paragraph("Just some prose with no table.")
    buffer = io.BytesIO()
    doc.save(buffer)
    response = course["fac"].upload(t(course, "/syllabus/import"), "scan.docx", buffer.getvalue())
    assert response.status_code == 201
    body = response.json()
    assert body["draft"]["outline"]["chapters"] == [] and "manually" in body["warnings"][0]
    assert body["draft"]["outline"]["course"]["name"] == "Intro to Entrepreneurship"
    counter = body["draft"]["counter"]
    assert course["fac"].delete(t(course, f"/syllabus/draft?expected_counter={counter}")).status_code == 204
    bad_type = course["fac"].upload(t(course, "/syllabus/import"), "x.exe", b"MZ")
    assert bad_type.status_code == 422


# ---------------- lessons: drafts, sanitising, targeting ----------------

def test_lesson_lifecycle_sanitises_and_keeps_published_visible(course):
    fac, st1 = course["fac"], course["st1"]
    item = new_item(course, "lesson", "Lesson 1")
    iid = item["id"]
    assert st1.get(learn(course, "/items")).json() == []                       # drafts are invisible
    assert st1.get(learn(course, f"/items/{iid}")).status_code == 404
    dirty = ('<p onclick="x()">Hello <strong>world</strong></p><script>alert(1)</script>'
             '<a href="javascript:alert(1)">bad</a><a href="https://example.com">ok</a>')
    done = edit_and_publish(course, iid, title="Lesson 1", body_html=dirty)
    assert done.status_code == 200, done.text
    lesson = st1.get(learn(course, f"/items/{iid}")).json()
    assert "<script" not in lesson["body_html"] and "onclick" not in lesson["body_html"]
    assert "javascript:" not in lesson["body_html"] and "<strong>world</strong>" in lesson["body_html"]
    assert 'rel="noopener noreferrer"' in lesson["body_html"]
    # editing: published stays visible while a new draft is changed
    draft = fac.post(t(course, f"/items/{iid}/draft")).json()["draft"]
    assert draft["version"] == 2 and "Hello" in draft["body_html"]
    fac.put(t(course, f"/items/{iid}/draft"), {"expected_counter": draft["counter"],
                                                "title": "Lesson 1 v2", "body_html": "<p>Second</p>"})
    assert st1.get(learn(course, f"/items/{iid}")).json()["title"] == "Lesson 1"
    done = fac.post(t(course, f"/items/{iid}/draft/publish"), {"expected_counter": 2})
    assert done.status_code == 200
    assert st1.get(learn(course, f"/items/{iid}")).json()["title"] == "Lesson 1 v2"
    # an item must be complete before publication
    empty = new_item(course, "lesson", "Empty")
    refused = edit_and_publish(course, empty["id"], title="Empty", body_html="")
    assert refused.status_code == 422 and refused.json()["error"]["code"] == "item_incomplete"


def test_section_targeting_and_archiving(course):
    a, b = course["sec_a"], course["sec_b"]
    item = new_item(course, "lesson", "Only for A", section_ids=[a["id"]])
    edit_and_publish(course, item["id"], title="Only for A", body_html="<p>A</p>")
    assert [i["title"] for i in course["st1"].get(learn(course, "/items")).json()] == ["Only for A"]
    assert course["st2"].get(learn(course, "/items")).json() == []
    assert course["st2"].get(learn(course, f"/items/{item['id']}")).status_code == 404
    other_term_section = course["fac"].patch(t(course, f"/items/{item['id']}"),
                                             {"section_ids": [str(uuid.uuid4())]})
    assert other_term_section.status_code == 422
    assert course["fac"].patch(t(course, f"/items/{item['id']}"),
                               {"section_ids": [a["id"], b["id"]]}).status_code == 200
    assert len(course["st2"].get(learn(course, "/items")).json()) == 1
    course["fac"].patch(t(course, f"/items/{item['id']}"), {"archived": True})
    assert course["st1"].get(learn(course, "/items")).json() == []


def test_never_published_items_can_be_deleted_but_published_ones_cannot(course):
    fac = course["fac"]
    draft_only = new_item(course, "lesson", "Scratch")
    assert fac.delete(t(course, f"/items/{draft_only['id']}")).status_code == 204
    assert fac.get(t(course, f"/items/{draft_only['id']}")).status_code == 404
    kept = new_item(course, "lesson", "Kept")
    edit_and_publish(course, kept["id"], title="Kept", body_html="<p>x</p>")
    assert fac.delete(t(course, f"/items/{kept['id']}")).status_code == 409
    # discarding a draft of a published item returns to the published revision
    draft = fac.post(t(course, f"/items/{kept['id']}/draft")).json()["draft"]
    assert fac.delete(t(course, f"/items/{kept['id']}/draft?expected_counter={draft['counter']}")).status_code == 204
    assert fac.get(t(course, f"/items/{kept['id']}")).json()["draft"] is None


def test_item_drafts_conflict_between_tabs(course):
    tab_b = Api("faculty@example.com")
    item = new_item(course, "lesson", "Shared")
    counter = course["fac"].get(t(course, f"/items/{item['id']}")).json()["draft"]["counter"]
    body = {"expected_counter": counter, "title": "Shared", "body_html": "<p>a</p>"}
    assert course["fac"].put(t(course, f"/items/{item['id']}/draft"), body).status_code == 200
    stale = tab_b.put(t(course, f"/items/{item['id']}/draft"), {**body, "body_html": "<p>b</p>"})
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "draft_revision_conflict"


# ---------------- files: authorised downloads ----------------

def test_file_download_authorisation(db, course):
    fac, st1, st2 = course["fac"], course["st1"], course["st2"]
    item = new_item(course, "file", "Reading", section_ids=[course["sec_a"]["id"]])
    iid = item["id"]
    counter = fac.get(t(course, f"/items/{iid}")).json()["draft"]["counter"]
    up = fac.upload(t(course, f"/items/{iid}/draft/file"), "reading.pdf", PDF, {"expected_counter": counter})
    assert up.status_code == 200, up.text
    file_id = up.json()["file"]["id"]
    # before publication: only the owning faculty can download
    assert fac.get(f"/api/files/{file_id}/download").status_code == 200
    assert st1.get(f"/api/files/{file_id}/download").status_code == 404
    edit_and_publish(course, iid, title="Reading")
    ok = st1.get(f"/api/files/{file_id}/download")
    assert ok.status_code == 200 and ok.content == PDF
    assert "attachment" in ok.headers["content-disposition"]
    # other section, other faculty, admin and anonymous callers are refused
    assert st2.get(f"/api/files/{file_id}/download").status_code == 404
    other = Api(make_account(db, "other@example.com", "faculty").email)
    assert other.get(f"/api/files/{file_id}/download").status_code == 404
    assert course["admin"].get(f"/api/files/{file_id}/download").status_code == 404
    assert TestClient(app).get(f"/api/files/{file_id}/download").status_code == 401
    # archived material is no longer downloadable by students
    fac.patch(t(course, f"/items/{iid}"), {"archived": True})
    assert st1.get(f"/api/files/{file_id}/download").status_code == 404
    # nothing is exposed under a guessable static path
    assert TestClient(app).get(f"/uploads/{file_id}").status_code == 404


def test_download_survives_new_drafts_and_republishing(course):
    fac, st1 = course["fac"], course["st1"]
    item = new_item(course, "file", "Reading")
    iid = item["id"]
    counter = fac.get(t(course, f"/items/{iid}")).json()["draft"]["counter"]
    file_id = fac.upload(t(course, f"/items/{iid}/draft/file"), "reading.pdf", PDF,
                         {"expected_counter": counter}).json()["file"]["id"]
    assert edit_and_publish(course, iid, title="Reading").status_code == 200
    assert st1.get(f"/api/files/{file_id}/download").status_code == 200
    # opening the editor creates a draft sharing the same file; students must still get it
    draft = fac.post(t(course, f"/items/{iid}/draft")).json()["draft"]
    assert draft["file"]["id"] == file_id
    assert st1.get(f"/api/files/{file_id}/download").status_code == 200
    # republish without changing the file: v1 superseded, v2 published, same file
    saved = fac.put(t(course, f"/items/{iid}/draft"), {"expected_counter": draft["counter"], "title": "Reading v2",
                                                        "reference_note": "Pages 1-3"})
    assert saved.status_code == 200
    assert fac.post(t(course, f"/items/{iid}/draft/publish"), {"expected_counter": saved.json()["counter"]}).status_code == 200
    assert st1.get(f"/api/files/{file_id}/download").status_code == 200
    assert fac.get(f"/api/files/{file_id}/download").status_code == 200
    # a file that only ever lived in a draft stays private
    other = new_item(course, "file", "Unpublished")
    c2 = fac.get(t(course, f"/items/{other['id']}")).json()["draft"]["counter"]
    private = fac.upload(t(course, f"/items/{other['id']}/draft/file"), "p.pdf", PDF, {"expected_counter": c2}).json()["file"]["id"]
    assert st1.get(f"/api/files/{private}/download").status_code == 404


def test_upload_validation_for_materials(course):
    fac = course["fac"]
    item = new_item(course, "file", "Reading")
    counter = fac.get(t(course, f"/items/{item['id']}")).json()["draft"]["counter"]
    url = t(course, f"/items/{item['id']}/draft/file")
    assert fac.upload(url, "run.exe", b"MZ", {"expected_counter": counter}).status_code == 422
    assert fac.upload(url, "fake.pdf", b"not a pdf", {"expected_counter": counter}).status_code == 422
    assert fac.upload(url, "fake.png", b"not a png", {"expected_counter": counter}).status_code == 422
    assert fac.upload(url, "ok.pdf", PDF, {"expected_counter": counter + 5}).status_code == 409
    lesson = new_item(course, "lesson", "A lesson")
    wrong_kind = fac.upload(t(course, f"/items/{lesson['id']}/draft/file"), "ok.pdf", PDF, {"expected_counter": 1})
    assert wrong_kind.status_code == 422
    unfinished = fac.post(t(course, f"/items/{item['id']}/draft/publish"), {"expected_counter": counter})
    assert unfinished.status_code == 422


def test_reference_drafts_save_freely_but_only_web_links_publish(course):
    fac = course["fac"]
    item = new_item(course, "reference", "Article")
    counter = fac.get(t(course, f"/items/{item['id']}")).json()["draft"]["counter"]
    half = fac.put(t(course, f"/items/{item['id']}/draft"),
                   {"expected_counter": counter, "title": "Article", "reference_url": "htt"})
    assert half.status_code == 200                                    # autosave never rejects half-typed input
    for bad in ("htt", "javascript:alert(1)", ""):
        saved = fac.put(t(course, f"/items/{item['id']}/draft"),
                        {"expected_counter": half.json()["counter"], "title": "Article", "reference_url": bad})
        refused = fac.post(t(course, f"/items/{item['id']}/draft/publish"),
                           {"expected_counter": saved.json()["counter"]})
        half = saved
        assert refused.status_code == 422 and refused.json()["error"]["code"] == "item_incomplete", bad
    assert course["st1"].get(learn(course, "/items")).json() == []
    done = edit_and_publish(course, item["id"], title="Article", reference_url="https://example.com/read",
                            reference_note="Chapter 2")
    assert done.status_code == 200
    shown = course["st1"].get(learn(course, f"/items/{item['id']}")).json()
    assert shown["reference_url"] == "https://example.com/read" and shown["reference_note"] == "Chapter 2"


def test_incomplete_grading_policy_saves_as_draft_but_cannot_publish(course):
    fac = course["fac"]
    outline, _, _ = make_outline()
    draft = fac.post(t(course, "/syllabus/draft")).json()
    policy = {"categories": [{"key": "quiz", "label": "", "weight": 100}],
              "periods": [{"key": "midterm", "label": "", "share": 100}]}
    saved = fac.put(t(course, "/syllabus/draft"), {"expected_counter": draft["counter"], "outline": outline,
                                                    "grading_policy": policy})
    assert saved.status_code == 200                                   # a cleared label is a normal editing state
    refused = fac.post(t(course, "/syllabus/draft/publish"), {"expected_counter": saved.json()["counter"]})
    assert refused.status_code == 422 and "Name every grading category" in refused.json()["error"]["message"]


def test_textless_pdf_syllabus_falls_back_to_manual_editing(course):
    from reportlab.pdfgen import canvas
    buffer = io.BytesIO()
    page = canvas.Canvas(buffer)
    page.showPage()                                                   # a valid PDF page with no text, like a scan
    page.save()
    response = course["fac"].upload(t(course, "/syllabus/import"), "scan.pdf", buffer.getvalue())
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["draft"]["outline"]["chapters"] == [] and "manually" in body["warnings"][0]
    assert body["draft"]["source_file"]["name"] == "scan.pdf"
    assert body["draft"]["outline"]["course"]["name"] == "Intro to Entrepreneurship"


# ---------------- anchors to the syllabus ----------------

def test_anchored_items_protect_syllabus_nodes(course):
    fac = course["fac"]
    outline, _, topic_id = make_outline()
    # an anchor must point at a real node of this subject's syllabus (draft or published)
    nowhere = fac.post(t(course, "/items"), {"kind": "lesson", "title": "x", "anchor_node_id": topic_id})
    assert nowhere.status_code == 422 and nowhere.json()["error"]["code"] == "anchor_invalid"
    draft = fac.post(t(course, "/syllabus/draft")).json()
    saved = fac.put(t(course, "/syllabus/draft"), {"expected_counter": draft["counter"], "outline": outline}).json()
    item = new_item(course, "lesson", "On topic", anchor_node_id=topic_id)       # exists in the draft
    # students cannot be shown a lesson whose topic is not in the published syllabus yet
    early = edit_and_publish(course, item["id"], title="On topic", body_html="<p>x</p>", anchor_node_id=topic_id)
    assert early.status_code == 409 and early.json()["error"]["code"] == "anchor_not_published"
    assert fac.post(t(course, "/syllabus/draft/publish"), {"expected_counter": saved["counter"]}).status_code == 200
    counter = fac.get(t(course, f"/items/{item['id']}")).json()["draft"]["counter"]
    ok = fac.post(t(course, f"/items/{item['id']}/draft/publish"), {"expected_counter": counter})
    assert ok.status_code == 200
    # removing the anchored topic in a new syllabus draft is refused until the item is resolved
    fresh = fac.post(t(course, "/syllabus/draft")).json()
    trimmed = {**fresh["outline"], "chapters": [{**fresh["outline"]["chapters"][0], "topics": []}]}
    saved = fac.put(t(course, "/syllabus/draft"), {"expected_counter": fresh["counter"], "outline": trimmed}).json()
    refused = fac.post(t(course, "/syllabus/draft/publish"), {"expected_counter": saved["counter"]})
    assert refused.status_code == 422 and "On topic" in refused.json()["error"]["message"]
    fac.patch(t(course, f"/items/{item['id']}"), {"archived": True})
    assert fac.post(t(course, "/syllabus/draft/publish"), {"expected_counter": saved["counter"]}).status_code == 200


def publish_lesson(c, title, **fields):
    item = new_item(c, "lesson", title, **fields)
    assert edit_and_publish(c, item["id"], title=title, body_html="<p>x</p>", **fields).status_code == 200
    return item["id"]


def test_lessons_follow_the_syllabus_then_faculty_order_and_students_see_progress(course):
    fac, st1 = course["fac"], course["st1"]
    outline, _, topic_id = make_outline()
    publish_syllabus(course, outline)
    loose = publish_lesson(course, "Loose")                        # no chapter: "Other materials", last
    second = publish_lesson(course, "Second", anchor_node_id=topic_id)
    first = publish_lesson(course, "First", anchor_node_id=topic_id)
    shown = st1.get(learn(course, "/items")).json()
    assert [i["title"] for i in shown] == ["Second", "First", "Loose"]       # topic before loose, then created order
    assert [i["up_next"] for i in shown] == [True, False, False] and shown[2]["anchor_node_id"] is None
    # faculty reorder the topic's group from the order they were looking at; a stale or foreign list is refused
    def order(expected, wanted, anchor=topic_id):
        return fac.put(t(course, "/items-order"), {"anchor_node_id": anchor, "expected_ids": expected, "item_ids": wanted})
    assert order([second, first], [first, second]).status_code == 204
    assert order([first], [first]).status_code == 409 and order([first, second, loose], [first, second, loose]).status_code == 409
    # a second tab that still shows the OLD order (same items, different order) must not overwrite the newer order
    stale = order([second, first], [second, first])
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "order_conflict"
    assert order([first, second], [first, first]).status_code == 409                       # not a permutation of the group
    assert [i["title"] for i in st1.get(learn(course, "/items")).json()] == ["First", "Second", "Loose"]
    assert st1.put(t(course, "/items-order"), {"expected_ids": [], "item_ids": []}).status_code == 403
    # an item published as unattached is not in the topic's group just because a draft of it names a topic
    draft = fac.post(t(course, f"/items/{loose}/draft")).json()["draft"]
    fac.put(t(course, f"/items/{loose}/draft"), {"expected_counter": draft["counter"], "title": "Loose", "body_html": "<p>x</p>", "anchor_node_id": topic_id})
    assert order([first, second], [second, first]).status_code == 204                      # the topic group is still just two items
    assert order([loose], [loose], anchor=None).status_code == 204                          # and Loose is still in "Other materials"
    assert order([second, first], [first, second]).status_code == 204
    # completion: marks the lesson, moves "up next", and each lesson names the one after it
    assert st1.get(learn(course, f"/items/{first}")).json()["next_lesson"]["title"] == "Second"
    st1.post(learn(course, f"/lessons/{first}/complete"))
    shown = st1.get(learn(course, "/items")).json()
    assert [i["completed"] for i in shown] == [True, False, False] and shown[1]["up_next"] is True
    assert st1.get(learn(course, f"/items/{first}")).json()["completed"] is True
    last = st1.get(learn(course, f"/items/{loose}")).json()
    assert last["next_lesson"] is None and last["lesson_number"] == 3 and last["lesson_total"] == 3
    # finishing only the LAST lesson is not "everything completed": the first unfinished one is offered
    st1.post(learn(course, f"/lessons/{loose}/complete"))
    last = st1.get(learn(course, f"/items/{loose}")).json()
    assert last["next_lesson"] is None and last["all_completed"] is False and last["first_incomplete"]["title"] == "Second"
    st1.post(learn(course, f"/lessons/{second}/complete"))
    assert st1.get(learn(course, f"/items/{loose}")).json()["all_completed"] is True
    # another student's completion is not mine
    assert [i["completed"] for i in course["st2"].get(learn(course, "/items")).json()] == [False] * 3


def test_announcements_load_in_bounded_pages_without_overlap(course):
    fac, st1 = course["fac"], course["st1"]
    for n in range(7):
        a = fac.post(t(course, "/announcements"), {"title": f"Notice {n}", "body": "x", "section_ids": []}).json()
        assert fac.post(t(course, f"/announcements/{a['id']}/publish")).status_code == 200
    seen, offset = [], 0
    while True:
        page = fac.get(t(course, f"/announcements?limit=3&offset={offset}")).json()
        seen += [a["title"] for a in page["items"]]
        if not page["has_more"]:
            break
        offset += 3
    assert len(seen) == 7 and len(set(seen)) == 7                                    # every one once, none repeated
    assert [len(st1.get(learn(course, f"/announcements?limit=3&offset={o}")).json()["items"]) for o in (0, 3, 6)] == [3, 3, 1]
    assert fac.get(t(course, "/announcements?limit=101")).status_code == 422          # the cap is enforced, never exceeded by the UI


# ---------------- coverage ----------------

def test_teaching_coverage_is_per_section_and_separate_from_students(course):
    outline, _, topic_id = make_outline()
    publish_syllabus(course, outline)
    fac, a = course["fac"], course["sec_a"]
    body = {"section_id": a["id"], "node_id": topic_id, "covered_on": "2026-07-01"}
    assert fac.put(t(course, "/coverage"), body).status_code == 204
    assert fac.put(t(course, "/coverage"), {**body, "node_id": uid()}).status_code == 422   # not in syllabus
    assert fac.put(t(course, "/coverage"), {**body, "section_id": str(uuid.uuid4())}).status_code == 422
    assert [c["node_id"] for c in course["st1"].get(learn(course, "/syllabus")).json()["covered"]] == [topic_id]
    assert course["st2"].get(learn(course, "/syllabus")).json()["covered"] == []   # other section
    assert fac.delete(t(course, f"/coverage?section_id={a['id']}&node_id={topic_id}")).status_code == 204
    assert course["st1"].get(learn(course, "/syllabus")).json()["covered"] == []


# ---------------- announcements ----------------

def test_announcements_are_targeted_and_locked_once_published(course):
    fac, a = course["fac"], course["sec_a"]
    created = fac.post(t(course, "/announcements"),
                       {"title": "Quiz Friday", "body": "Bring a pen.", "section_ids": [a["id"]]})
    assert created.status_code == 201
    aid = created.json()["id"]
    assert course["st1"].get(learn(course, "/announcements")).json()["items"] == []          # draft: hidden
    edited = fac.put(t(course, f"/announcements/{aid}"),
                     {"title": "Quiz Friday!", "body": "Bring a pen.", "section_ids": [a["id"]]})
    assert edited.status_code == 200
    assert fac.post(t(course, f"/announcements/{aid}/publish")).status_code == 200
    assert [x["title"] for x in course["st1"].get(learn(course, "/announcements")).json()["items"]] == ["Quiz Friday!"]
    assert course["st2"].get(learn(course, "/announcements")).json()["items"] == []          # other section
    locked = fac.put(t(course, f"/announcements/{aid}"), {"title": "x", "body": "y", "section_ids": []})
    assert locked.status_code == 409
    assert fac.post(t(course, f"/announcements/{aid}/archive")).status_code == 200
    assert course["st1"].get(learn(course, "/announcements")).json()["items"] == []


# ---------------- closed terms and withdrawn students ----------------

def test_closed_term_blocks_every_teaching_write_but_not_reads(course):
    fac = course["fac"]
    item = new_item(course, "lesson", "Before close")
    counter = fac.get(t(course, f"/items/{item['id']}")).json()["draft"]["counter"]
    publish_syllabus(course)
    term_id = course["term"]["id"]
    assert course["admin"].post(f"/api/terms/{term_id}/close").status_code == 200
    closed = lambda r: r.status_code == 409 and r.json()["error"]["code"] == "term_closed"
    assert closed(fac.post(t(course, "/syllabus/draft")))
    assert closed(fac.post(t(course, "/items"), {"kind": "lesson", "title": "x"}))
    assert closed(fac.put(t(course, f"/items/{item['id']}/draft"),
                          {"expected_counter": counter, "title": "autosave", "body_html": "<p>x</p>"}))
    assert closed(fac.post(t(course, "/announcements"), {"title": "x", "body": "y", "section_ids": []}))
    assert fac.get(t(course, "/syllabus")).status_code == 200
    assert fac.get(t(course, "/items")).status_code == 200
    assert course["st1"].get(learn(course, "/syllabus")).json()["published"]["version"] == 1
    assert course["admin"].post(f"/api/terms/{term_id}/reopen", {"reason": "Fix typo"}).status_code == 200
    assert fac.post(t(course, "/syllabus/draft")).status_code == 200


def test_withdrawn_student_gets_no_course_content(course):
    publish_syllabus(course)
    item = new_item(course, "lesson", "Lesson")
    edit_and_publish(course, item["id"], title="Lesson", body_html="<p>x</p>")
    st1 = course["st1"]
    assert st1.get(learn(course, f"/items/{item['id']}")).status_code == 200
    s1 = course["students"][0]
    course["admin"].delete(f"/api/sections/{course['sec_a']['id']}/members/{s1.id}?reason=Left%20the%20programme")
    for path in ("/syllabus", "/items", f"/items/{item['id']}", "/announcements"):
        assert st1.get(learn(course, path)).status_code == 404, path
    # the subject is still listed (labelled withdrawn) so the student keeps their history
    subjects = st1.get("/api/me/subjects").json()
    assert [s["enrollment_status"] for s in subjects] == ["withdrawn"]

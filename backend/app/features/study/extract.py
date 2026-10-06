"""Readable text from published teaching content: lesson HTML, PDF, DOCX and PPTX files.

Scanned documents and images have no text layer (there is no OCR): they simply produce no chunks,
which the faculty page reports so the material is not silently assumed to be searchable."""
import html
import io
import re
from pathlib import Path

import nh3
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

MAX_CHARS = 300_000
MAX_PAGES = 400
BLOCK_END = re.compile(r"</(p|div|li|tr|h[1-6]|blockquote|pre|table)>|<br\s*/?>", re.IGNORECASE)


def html_to_text(markup):
    """Tags stripped, paragraph breaks kept. The input is already server-sanitised HTML."""
    spaced = BLOCK_END.sub("\n", markup or "")
    text = html.unescape(nh3.clean(spaced, tags=set()))
    return re.sub(r"[ \t]+", " ", re.sub(r"\n\s*\n+", "\n\n", text)).strip()


def pdf_sections(data):
    reader = PdfReader(io.BytesIO(data))
    out = []
    for number, page in enumerate(reader.pages[:MAX_PAGES], start=1):
        text = (page.extract_text() or "").strip()
        if text:
            out.append((f"Page {number}", text))
    return out


def docx_sections(data):
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                if cell.text.strip() and cell.text.strip() not in cells:
                    cells.append(cell.text.strip())
            if cells:
                parts.append(" | ".join(cells))
    return [("Document", "\n".join(parts))] if parts else []


def pptx_sections(data):
    deck = Presentation(io.BytesIO(data))
    out = []
    for number, slide in enumerate(deck.slides, start=1):
        text = "\n".join(shape.text_frame.text for shape in slide.shapes
                         if shape.has_text_frame and shape.text_frame.text.strip())
        if text.strip():
            out.append((f"Slide {number}", text))
    return out


READERS = {".pdf": pdf_sections, ".docx": docx_sections, ".pptx": pptx_sections}


def file_sections(path, original_name):
    reader = READERS.get(Path(original_name).suffix.lower())
    if reader is None:
        return []
    sections, total = [], 0
    for locator, text in reader(Path(path).read_bytes()):
        total += len(text)
        sections.append((locator, text[:MAX_CHARS]))
        if total >= MAX_CHARS:
            break
    return sections

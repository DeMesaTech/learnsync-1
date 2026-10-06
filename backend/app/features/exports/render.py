"""Real XLSX and PDF files from a grade table."""
import re
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value):
    """A spreadsheet runs text that starts with = + - @ as a formula: neutralise names and notes."""
    if isinstance(value, str) and value.startswith(FORMULA_START):
        return "'" + value
    return value


def safe_filename(name):
    return re.sub(r"[^A-Za-z0-9._-]+", "-", name).strip("-.")[:80] or "export"


def to_xlsx(table):
    book = Workbook()
    sheet = book.active
    sheet.title = "Grades"
    row = 1
    if table["banner"]:
        cell = sheet.cell(row=row, column=1, value=safe_cell(table["banner"]))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="B42318")
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(len(table["columns"]), 4))
        row += 1
    sheet.cell(row=row, column=1, value=safe_cell(table["title"])).font = Font(bold=True, size=14)
    row += 1
    for label, value in table["meta"]:
        sheet.cell(row=row, column=1, value=safe_cell(label)).font = Font(bold=True)
        sheet.cell(row=row, column=2, value=safe_cell(value))
        row += 1
    for note in table["notes"]:
        sheet.cell(row=row, column=1, value=safe_cell(note)).font = Font(italic=True)
        row += 1
    row += 1
    header = row
    for col, name in enumerate(table["columns"], start=1):
        cell = sheet.cell(row=header, column=col, value=safe_cell(name))
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DCE6F2")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r, data in enumerate(table["rows"], start=header + 1):
        for col, value in enumerate(data, start=1):
            cell = sheet.cell(row=r, column=col, value=safe_cell(value))
            if isinstance(value, float):
                cell.number_format = "0.00"
    sheet.freeze_panes = sheet.cell(row=header + 1, column=4 if table["columns"][0] == "Rank" else 3)
    for col in range(1, len(table["columns"]) + 1):
        widest = max([len(str(table["columns"][col - 1]))] +
                     [len(str(d[col - 1])) for d in table["rows"] if d[col - 1] is not None])
        sheet.column_dimensions[get_column_letter(col)].width = min(max(widest + 2, 10), 60)
    sheet.page_setup.orientation = "landscape"
    sheet.print_title_rows = f"{header}:{header}"
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def column_widths(count, ranked=False):
    """Column widths (points) for a landscape A4 table. Many category columns: the fixed columns shrink,
    the name column gets a larger share than a percentage, so names stay readable. `ranked` means a Rank
    column was put in front of the usual layout."""
    last = count - 1
    lead = 1 if ranked else 0
    available = landscape(A4)[0] - 20 * mm
    crowded = count - lead > 11
    base = ({0: 20 * mm, 2: 16 * mm, -3: 13 * mm, -2: 15 * mm, -1: 46 * mm} if crowded
            else {0: 22 * mm, 2: 20 * mm, -3: 15 * mm, -2: 16 * mm, -1: 60 * mm})
    fixed = {(k + lead if k >= 0 else last + 1 + k): v for k, v in base.items()}
    if ranked:
        fixed[0] = 11 * mm
    weights = {i: (2.6 if i == 1 + lead else 1.0) for i in range(count) if i not in fixed}
    unit = (available - sum(fixed.values())) / sum(weights.values())
    return [fixed.get(i, unit * weights.get(i, 1.0)) for i in range(count)]


def to_pdf(table):
    styles = getSampleStyleSheet()
    size = 6.3 if len(table["columns"]) > 11 else 7.5
    small = styles["BodyText"].clone("small", fontSize=size, leading=size + 1.5)
    head = styles["BodyText"].clone("head", fontSize=size, leading=size + 1.5, textColor=colors.white,
                                    fontName="Helvetica-Bold")
    banner = table["banner"]

    def decorate(canvas, doc):
        canvas.saveState()
        width, height = landscape(A4)
        if banner:   # on EVERY page, so a printed preview cannot be mistaken for released grades
            canvas.setFillColor(colors.HexColor("#B42318"))
            canvas.rect(10 * mm, height - 13 * mm, width - 20 * mm, 7 * mm, stroke=0, fill=1)
            canvas.setFillColor(colors.white)
            canvas.setFont("Helvetica-Bold", 9)
            canvas.drawString(12 * mm, height - 11.2 * mm, banner)
        canvas.setFillColor(colors.black)
        canvas.setFont("Helvetica", 7)
        canvas.drawString(10 * mm, 6 * mm, table["notes"][0])
        canvas.drawRightString(width - 10 * mm, 6 * mm, f"Page {doc.page}")
        canvas.restoreState()

    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=17 * mm, bottomMargin=12 * mm, title=table["title"],
                            author="LearnSync")
    story = [Paragraph(escape(table["title"]), styles["Title"])]
    for label, value in table["meta"]:
        story.append(Paragraph(f"<b>{escape(label)}:</b> {escape(str(value))}", small))
    for note in table["notes"][1:]:
        story.append(Paragraph(f"<i>{escape(note)}</i>", small))
    story.append(Spacer(1, 4 * mm))
    data = [[Paragraph(escape(c), head) for c in table["columns"]]]
    for row in table["rows"]:
        data.append([Paragraph(escape("" if v is None else (f"{v:.2f}" if isinstance(v, float) else str(v))), small)
                     for v in row])
    widths = column_widths(len(table["columns"]), table["columns"][0] == "Rank")
    grid = Table(data, colWidths=widths, repeatRows=1)
    grid.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#9AA5B1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6FA")])]))
    story.append(grid)
    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return out.getvalue()

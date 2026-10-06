"""Best-effort syllabus extraction. Output is only ever a draft that faculty reviews/edits.

Tuned on the institution's DOCX syllabus template (course-information table, learning outcomes,
week-by-week coverage table, evaluation, references, requirements); anything else degrades to
"nothing recognised" and the caller falls back to manual editing."""
import io
import re
import uuid

from docx import Document
from pypdf import PdfReader

from .schemas import Chapter, CourseInfo, Outcome, Outline, Topic

WEEK = re.compile(r"Week\s*\d+(?:\s*[–—-]\s*\d+)?", re.IGNORECASE)
CHAPTER = re.compile(r"^\s*(chapter|unit|module)\s+\d+\b", re.IGNORECASE)
EXAM = re.compile(r"\b(midterm|mid-term|final)\s+exam", re.IGNORECASE)
INFO_LABELS = {
    "course no": "code", "course name": "name", "course description": "description",
    "credit units": "units", "contact hours": "contact_hours", "pre-requisite": "prerequisite",
    "values integration": "values",
}


CAPS = {"code": 60, "name": 300, "description": 5000, "units": 10, "contact_hours": 60,
        "prerequisite": 300, "values": 1000, "references": 10000, "requirements": 5000}


def new_id():
    return uuid.uuid4()


def clean(text):
    return re.sub(r"[ \t]+", " ", text.replace("\xa0", " ")).strip()


def lines_of(text):
    return [ln for ln in (clean(x).lstrip("•-– ").strip() for x in text.splitlines()) if ln]


def unique_cells(row):
    """Merged cells repeat in python-docx; keep each once and drop empty/bullet-only cells."""
    seen, out = set(), []
    for cell in row.cells:
        if id(cell._tc) in seen:
            continue
        seen.add(id(cell._tc))
        if re.search(r"[A-Za-z0-9]", cell.text):
            out.append(cell.text)
    return out


def parse_docx(path_bytes):
    doc = Document(io.BytesIO(path_bytes))
    info, outcomes, chapters = {}, [], []
    for table in doc.tables:
        for row in table.rows:
            cells = unique_cells(row)
            if not cells:
                continue
            head = clean(cells[0])
            label = re.sub(r"^\d+\.?\s*", "", head).rstrip(".: ").lower()
            if len(cells) == 2 and label in INFO_LABELS:
                info[INFO_LABELS[label]] = clean(cells[1])
            elif head.startswith("•") and len(cells) >= 1:
                outcomes.append(lines_of(head)[0] if lines_of(head) else "")
            elif label.startswith("references"):
                info["references"] = "\n".join(lines_of(cells[0])[1:])
            elif label.startswith("classroom requirements"):
                info["requirements"] = "\n".join(lines_of(cells[0])[1:])
            elif label.startswith("examination") and "%" in head:
                info.setdefault("exam_evaluation", "\n".join(lines_of(cells[0])))
            elif WEEK.search(head) and len(cells) >= 2:
                chapters.append(coverage_row(cells))
    evaluation = []
    capture = False
    for p in doc.paragraphs:
        text = clean(p.text)
        if re.match(r"^\d+\.?\s*Course Evaluation", text, re.IGNORECASE):
            capture = True
        elif capture and text:
            evaluation.append(text)
    if info.pop("exam_evaluation", ""):
        evaluation.append("Examination: see syllabus table")
    exam_text = [clean(t) for t in evaluation if t]
    course = CourseInfo(**{k: v[:CourseInfo.model_fields[k].metadata[0].max_length]
                           for k, v in info.items() if k in CourseInfo.model_fields},
                        evaluation_note="\n".join(exam_text)[:3000])
    return Outline(course=course,
                   outcomes=[Outcome(id=new_id(), text=t[:1000]) for t in outcomes if t],
                   chapters=[c for c in chapters if c])


def coverage_row(cells):
    """[week, ilo, topics, activities, assessment] -> Chapter (or an exam row)."""
    week = clean(WEEK.search(cells[0]).group(0)) if WEEK.search(cells[0]) else ""
    rest = cells[1:]
    if len(rest) == 1 and EXAM.search(rest[0]):
        return Chapter(id=new_id(), kind="exam", title=clean(rest[0]).title(), weeks=week)
    if len(rest) < 2:
        return None
    ilo, topics = rest[0], rest[1]
    activities = rest[2] if len(rest) > 2 else ""
    assessment = rest[-1] if len(rest) > 3 else ""
    lines = lines_of(topics)
    if not lines:
        return None
    title, topic_lines = lines[0], lines[1:]
    if not CHAPTER.match(title) and len(lines) > 1 and CHAPTER.match(lines[1]):
        title, topic_lines = lines[1], lines[2:]
    return Chapter(id=new_id(), title=title[:300], weeks=week,
                   ilo=" ".join(lines_of(ilo))[:3000],
                   activities=", ".join(lines_of(activities))[:3000],
                   assessment=", ".join(lines_of(assessment))[:3000],
                   topics=[Topic(id=new_id(), title=t[:300]) for t in topic_lines[:100]])


def parse_pdf(data):
    """PDFs are only lightly read: chapter headings become chapters; faculty fills the rest."""
    lines = [clean(ln) for page in PdfReader(io.BytesIO(data)).pages
             for ln in (page.extract_text() or "").splitlines()]
    chapters = [Chapter(id=new_id(), title=ln[:300]) for ln in lines if CHAPTER.match(ln)]
    return Outline(chapters=chapters[:100])


def parse_syllabus(path, ext):
    data = path.read_bytes()
    outline = parse_docx(data) if ext == ".docx" else parse_pdf(data)
    warnings = []
    if outline.chapters or outline.outcomes:
        warnings.append("Review everything: extraction is a starting point, not authoritative.")
    if not outline.course.name:
        warnings.append("The course name was not recognised.")
    return outline, warnings
